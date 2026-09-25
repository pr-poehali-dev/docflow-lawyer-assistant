"""Рендеринг блоков документа в DOCX и PDF."""
import base64
import io

from docx import Document as DocxDocument
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from fonts_data import REGULAR_B64, BOLD_B64, ITALIC_B64

ALIGN_MAP_DOCX = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
}


def render_docx(blocks, title):
    doc = DocxDocument()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)

    for b in blocks:
        p = doc.add_paragraph()
        p.alignment = ALIGN_MAP_DOCX.get(b["align"], WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.space_after = Pt(b["spacing_after"])
        p.paragraph_format.line_spacing = 1.3
        run = p.add_run(b["text"])
        run.bold = b["bold"]
        run.italic = b.get("italic", False)
        run.font.size = Pt(b["size"])

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


_FONTS_REGISTERED = False


def _register_fonts():
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return
    pdfmetrics.registerFont(TTFont('Body', io.BytesIO(base64.b64decode(REGULAR_B64))))
    pdfmetrics.registerFont(TTFont('Body-Bold', io.BytesIO(base64.b64decode(BOLD_B64))))
    pdfmetrics.registerFont(TTFont('Body-Italic', io.BytesIO(base64.b64decode(ITALIC_B64))))
    _FONTS_REGISTERED = True


def _wrap_text(c, text, font_name, font_size, max_width):
    words = text.split(' ')
    lines = []
    current = ""
    for w in words:
        trial = (current + " " + w).strip()
        if c.stringWidth(trial, font_name, font_size) <= max_width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = w
    if current:
        lines.append(current)
    return lines or [""]


def render_pdf(blocks, title):
    _register_fonts()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    left_margin = 2.5 * cm
    right_margin = 2 * cm
    top_margin = 2 * cm
    bottom_margin = 2 * cm
    max_width = width - left_margin - right_margin

    y = height - top_margin

    def new_page():
        nonlocal y
        c.showPage()
        y = height - top_margin

    for b in blocks:
        font_name = 'Body-Bold' if b["bold"] else ('Body-Italic' if b.get('italic') else 'Body')
        font_size = b["size"]
        line_height = font_size * 1.4

        lines = _wrap_text(c, b["text"], font_name, font_size, max_width)
        for line in lines:
            if y < bottom_margin:
                new_page()
            c.setFont(font_name, font_size)
            text_width = c.stringWidth(line, font_name, font_size)
            if b["align"] == "center":
                x = left_margin + (max_width - text_width) / 2
            elif b["align"] == "right":
                x = left_margin + max_width - text_width
            else:
                x = left_margin
            c.drawString(x, y, line)
            y -= line_height
        y -= b["spacing_after"] * 0.6

    c.save()
    buf.seek(0)
    return buf.read()